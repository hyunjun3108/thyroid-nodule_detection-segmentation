#!/usr/bin/env python
"""석사 졸업논문 연구계획서(YOLOv11 갑상선 결절 detection+segmentation) -> .docx 생성."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = "/data/HJ_thyroid/YOLOv11_갑상선결절_연구계획서.docx"

doc = Document()
# 기본 폰트
style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style.font.size = Pt(10.5)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)
    return p


def para(text, bold=False, italic=False, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p


def bullet(text):
    doc.add_paragraph(text, style="List Bullet")


def numbered(text):
    doc.add_paragraph(text, style="List Number")


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(t.rows[0].cells):
        hd.text = ""
        run = hd.paragraphs[0].add_run(headers[i])
        run.bold = True
        run.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


# ===== 표지 =====
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("YOLOv11 기반 갑상선 결절 초음파 영상의\n검출 및 분할 연구")
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(0x1F, 0x3B, 0x63)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run("― 경계 정교화 모듈과 경량화를 중심으로 ―")
rs.italic = True
rs.font.size = Pt(13)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("석사학위 논문 연구계획서").font.size = Pt(11)
doc.add_paragraph()

# ===== 1. 연구 배경 =====
h("1. 연구 배경 및 필요성")
para("갑상선 결절은 성인에서 매우 흔하게 발견되며, 초음파 검사는 결절의 존재·위치·경계·내부 특성을 "
     "평가하는 일차 진단 도구이다. 그러나 초음파 영상은 스페클 잡음(speckle noise), 낮은 대조도, "
     "탐촉자·기기·검사자에 따른 편차가 크고, 결절 경계가 흐릿하여 판독자 간 변이가 크다는 한계가 있다. "
     "이에 따라 결절을 자동으로 검출(detection)하고 정밀하게 분할(segmentation)하는 딥러닝 기법의 필요성이 크다.")
para("본 연구의 문제의식은 다음 세 가지이다.", bold=True)
bullet("① 경계 모호성: 초음파 결절의 경계가 불명확하여 기존 분할 모델은 경계 부근에서 오차가 크다. "
       "→ 경계 품질을 직접 개선하는 모듈·손실이 필요하다.")
bullet("② 임상 배포 제약: 실제 장비·엣지 환경에 적용하려면 정확도를 유지하면서도 파라미터·연산량을 줄인 "
       "경량 모델이 필요하다.")
bullet("③ 일반화 한계: 단일 데이터셋 학습 모델은 다른 기관·기기 영상에서 성능이 급락한다. "
       "→ 여러 공개 데이터셋을 혼합하여 도메인 일반화를 확보해야 한다.")

# ===== 2. 연구 목적 및 기여 =====
h("2. 연구 목적 및 기여")
para("본 연구의 목적은 YOLOv11 분할 모델(YOLOv11-seg)을 기반으로 갑상선 결절의 검출과 분할을 하나의 "
     "멀티태스크 모델로 동시에 수행하고, 경계 정교화와 경량화를 통해 정확도와 효율을 모두 확보하는 것이다.")
para("주요 기여(Contributions)", bold=True)
numbered("경계 정교화 모듈(Boundary Refinement Module, BRM) 및 경계 인식 손실(boundary-aware loss)을 "
         "제안하여 흐릿한 결절 경계에서의 검출·분할 정확도를 향상한다. (모듈+손실 이중 기여)")
numbered("경량 백본/블록 교체 기반 경량화를 수행하여, 정확도 저하를 최소화하면서 파라미터·연산량·추론속도를 "
         "개선한다. (정확도–효율 trade-off 분석)")
numbered("다중 공개 데이터셋(TN3K, DDTI 등) 혼합 학습과 Leave-One-Dataset-Out(LODO) 평가를 통해 "
         "도메인 일반화 성능을 정량적으로 검증한다.")

# ===== 3. 관련 연구 =====
h("3. 관련 연구")
bullet("객체 검출·분할: YOLO 계열(실시간 검출/인스턴스 분할), Mask 기반 분할 모델.")
bullet("갑상선 초음파 분할: TRFE-Net(갑상선 영역 사전지식 기반), TNSCUI2020 챌린지 계열 캐스케이드 분할.")
bullet("경계 인식 기법: Boundary IoU, Active Contour Loss, Boundary Loss(거리맵 기반), 경계 주의(attention) 모듈.")
bullet("경량화: GhostNet/MobileNet 계열 경량 블록, 채널 프루닝, 지식 증류(KD), INT8 양자화/TensorRT.")

# ===== 4. 데이터셋 =====
h("4. 데이터셋")
para("일반화를 위해 마스크(픽셀 단위 정답)가 포함된 공개 갑상선 초음파 분할 데이터셋을 혼합 사용한다. "
     "모든 결절을 단일 클래스(nodule)로 통일하며, 이진 마스크를 임계값 이진화(>127) 후 컨투어 기반으로 "
     "YOLO-seg 폴리곤 라벨로 변환한다.")
table(["데이터셋", "구성/규모", "라벨", "용도"],
      [["TN3K", "trainval 2,879 + test 614 (총 3,493)", "결절 마스크(JPG)", "핵심 학습/평가"],
       ["DDTI", "image + mask (PNG)", "결절 마스크", "도메인 다양성/LODO 일반화"],
       ["TG3K", "3,585 (갑상선 gland)", "gland 마스크", "보조(사전학습) - 결절 아님, 선택"]],
      widths=[3.0, 6.5, 3.5, 3.5])
para("※ 데이터셋 간 해상도·기기·라벨 기준 편차가 존재하므로, 전처리(그레이스케일→3채널, imgsz=640 정규화)를 "
     "통일하고 손상·무주석 파일을 제거한다. 양성/악성 분류 라벨(DDTI category)은 본 연구 범위 밖이며 향후 확장 과제로 둔다.",
     italic=True, size=9.5)

# ===== 5. 제안 방법 =====
h("5. 제안 방법")
h("5.1 베이스라인: YOLOv11-seg 멀티태스크", level=2)
para("YOLOv11-seg의 Segment head는 바운딩 박스와 마스크를 동시에 출력하므로, 하나의 모델로 detection과 "
     "segmentation을 함께 수행한다. yolo11n-seg / yolo11s-seg를 기준 모델로 사용한다.")
h("5.2 경계 정교화 모듈(BRM) + 경계 손실", level=2)
bullet("모듈(BRM): 마스크 분기의 프로토타입(Proto) 경로에 경량 정교화 모듈을 삽입한다. 학습형 에지(Sobel) "
       "특징과 경계 영역 강조 attention을 프로토타입에 게이트 융합(residual)하여 경계 표현을 강화한다.")
bullet("손실: 분할 손실에 경계 항을 추가한다. 정답 마스크의 모폴로지 그래디언트로 얻은 경계 밴드에 대한 "
       "가중 Dice/BCE를 기본으로 하고, 필요 시 Active Contour/Boundary(거리맵) 손실과 비교한다.")
bullet("설계 유의점: 과거 유사 프로젝트에서 대형 attention 모듈이 오히려 성능을 저하시킨 사례가 있으므로, "
       "삽입 위치(백본/Proto/head)와 손실 가중치를 ablation으로 신중히 검증한다.")
h("5.3 경량화: 경량 백본/블록 교체", level=2)
bullet("백본/넥의 Conv·C3k2 블록을 GhostConv/Ghost-C3k2 또는 MobileNet/EfficientViT 계열 경량 블록으로 "
       "치환하여 파라미터·연산량을 절감한다. 경계 정교화 모듈(BRM)은 유지한다.")
bullet("(선택) 대형 BRM 모델을 teacher로 한 지식 증류로 경량 student의 정확도를 회복한다.")
bullet("INT8 양자화 및 TensorRT/ONNX export로 배포 최적화 후 GPU·CPU 추론속도를 측정한다.")

# ===== 6. 실험 설계 =====
h("6. 실험 설계 및 평가")
para("데이터 분할", bold=True)
bullet("Pooled 5-fold 교차검증: 전체 혼합 데이터를 데이터셋 출처 기준으로 층화하여 5-fold로 평가(주 성능표).")
bullet("Leave-One-Dataset-Out(LODO): 한 데이터셋으로 학습, 미학습 데이터셋으로 평가하여 도메인 일반화 정량화.")
para("평가지표", bold=True)
table(["범주", "지표"],
      [["검출(Detection)", "box mAP@50, mAP@50-95, Precision, Recall"],
       ["분할(Segmentation)", "mask mAP@50, mAP@50-95, Dice, IoU"],
       ["경계(Boundary)", "Boundary IoU, HD95(95% Hausdorff), ASSD"],
       ["효율(Efficiency)", "#Params, GFLOPs, FPS(GPU/CPU), 모델 크기"]],
      widths=[4.0, 12.0])
para("Ablation 구성 (경계 기여 검증)", bold=True)
table(["구성", "설명"],
      [["Baseline", "YOLOv11-seg (수정 없음)"],
       ["+ BRM", "경계 정교화 모듈만 추가"],
       ["+ Boundary Loss", "경계 손실만 추가"],
       ["+ BRM + Loss (제안)", "모듈+손실 결합"],
       ["+ 경량화", "제안 모델의 경량 백본/블록 버전"]],
      widths=[5.0, 11.0])

# ===== 7. 추진 일정 =====
h("7. 단계별 추진 일정")
table(["단계", "내용", "기간(주)", "산출물"],
      [["Phase 0", "환경·파이프라인 검증(격리 env, 포크 설치, seg 동작 확인)", "0.5", "동작 검증 로그"],
       ["Phase 1", "공개 데이터 수집·polygon 변환·분할 구성(Pooled/LODO)", "1.5–2", "통합 데이터셋, data.yaml"],
       ["Phase 2", "YOLOv11-seg 베이스라인 학습·평가(경계지표 코드 추가)", "1", "베이스라인 성능표"],
       ["Phase 3", "경계 정교화 모듈+경계 손실 구현·ablation", "2.5–3", "제안 모델, ablation 표"],
       ["Phase 4", "경량 백본/블록 교체·양자화·효율 벤치", "2–3", "경량 모델, Pareto 표"],
       ["Phase 5", "최종 종합 평가·정성 결과·논문 산출물 정리", "1", "논문 표·그림"]],
      widths=[2.2, 8.3, 2.0, 3.5])
para("※ Phase 0은 이미 완료(격리 env thyroid_seg 구성, YOLOv11-seg 파이프라인 mask 지표 생성 검증).",
     italic=True, size=9.5)

# ===== 8. 기대효과 =====
h("8. 기대효과 및 활용")
bullet("경계 정교화를 통한 결절 분할 정확도 향상 → 크기·경계 계측 신뢰도 개선(임상 판독 보조).")
bullet("경량화를 통한 실시간·엣지 배포 가능성 확보(휴대형 초음파, 검진 현장).")
bullet("다중 데이터셋 일반화 검증 → 실제 임상 환경 적용 시 강건성 근거 제시.")

# ===== 9. 위험요소 =====
h("9. 위험요소 및 대응")
table(["위험요소", "대응 방안"],
      [["공개 데이터 라이선스/재배포 제약", "논문에는 지표·방법만 공개, 원본 재배포 금지 준수"],
       ["데이터셋 간 도메인 갭(LODO 성능 저하)", "혼합 학습·증강, 성능 저하를 일반화 논의로 분석"],
       ["신규 모듈이 성능을 저해할 위험(전례 존재)", "위치·가중치 ablation 필수, 동일 조건 baseline 비교"],
       ["경량화 시 정확도 하락", "지식 증류로 회복, 정확도–효율 Pareto로 절충점 제시"]],
      widths=[6.5, 9.5])

# ===== 10. 개발 환경 =====
h("10. 개발 환경 및 인프라")
bullet("하드웨어: NVIDIA A100 80GB × 2–3.")
bullet("소프트웨어: Python 3.10, PyTorch 2.5.1(CUDA 12.1), Ultralytics(YOLOv11) 커스텀 포크(편집형 설치).")
bullet("격리 환경: 갑상선 전용 conda env(thyroid_seg) 및 프로젝트 디렉터리(/data/HJ_thyroid).")
bullet("배포 최적화: TensorRT(INT8)·ONNX export 환경.")

doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = foot.add_run("본 계획서는 연구 진행에 따라 실험 결과를 반영하여 갱신될 수 있음.")
fr.italic = True
fr.font.size = Pt(9)

cp = doc.core_properties                # 라이브러리 기본 문서 속성 정리
cp.author = cp.last_modified_by = cp.comments = cp.keywords = cp.category = ""
doc.save(OUT)
print("생성 완료:", OUT)
