#!/usr/bin/env python
"""
석사논문 발표자료(.pptx) 생성 — 배경/문제정의 · 선행연구/가설 · 제안방법 · 실험결과.

runs/figures/*.png 을 그대로 삽입한다. 수치는 모두 실측값
(runs/phase4_quant/final_table.csv, runs/stats/paired_stats.csv 등)에서 가져온 것.

사용:
  /data/my_envs/thyroid_seg/bin/python scripts/make_ppt.py
출력: /data/HJ_thyroid/갑상선결절_경계정교화_경량화_발표.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches as In, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image

ROOT = "/data/HJ_thyroid"
FIG = f"{ROOT}/runs/figures"
OUT = f"{ROOT}/갑상선결절_경계정교화_경량화_발표.pptx"

# 본문 팔레트 — 그림(make_figures.py)과 동일 계열로 맞춤
INK = RGBColor(0x1A, 0x1A, 0x19)
MUTED = RGBColor(0x5C, 0x5C, 0x5A)
FAINT = RGBColor(0x8E, 0x8E, 0x8B)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
RULE = RGBColor(0xD8, 0xD8, 0xD5)
PANEL = RGBColor(0xF5, 0xF4, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "맑은 고딕"
W, H = 13.333, 7.5
ML, MR = 0.68, 0.68          # 좌우 여백
TOP_TITLE, TOP_BODY = 0.46, 1.42
BOT = 6.86


# ----------------------------------------------------------------- 저수준 유틸
def set_font(run, size=14, bold=False, color=INK, font=FONT, italic=False):
    f = run.font
    f.size, f.bold, f.italic = Pt(size), bold, italic
    f.color.rgb = color
    f.name = font
    # 한글은 latin 이 아니라 동아시아 글꼴 슬롯을 봐야 한다
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        e = rPr.find(qn(tag))
        if e is None:
            e = rPr.makeelement(qn(tag), {})
            rPr.append(e)
        e.set("typeface", font)


def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def para(tf, text, size=14, bold=False, color=INK, space_before=0, space_after=6,
         align=PP_ALIGN.LEFT, first=False, line=1.28, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_before, p.space_after = Pt(space_before), Pt(space_after)
    p.line_spacing = line
    set_font(p.add_run(), size)          # 빈 run 방지용으로 즉시 텍스트 지정
    p.runs[0].text = text
    set_font(p.runs[0], size, bold, color, italic=italic)
    return p


def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, In(x), In(y), In(w), In(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    return sh


def picture(slide, path, x, y, w, h):
    """(x,y,w,h) 박스 안에 비율 유지하며 중앙 배치."""
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    return slide.shapes.add_picture(path, In(x + (w - pw) / 2), In(y + (h - ph) / 2),
                                    In(pw), In(ph))


def cell_border(cell, edges=("B",), color="D8D8D5", w=9525):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"L": "a:lnL", "R": "a:lnR", "T": "a:lnT", "B": "a:lnB"}
    for e in edges:
        ln = tcPr.makeelement(qn(tag[e]), {"w": str(w), "cap": "flat",
                                           "cmpd": "sng", "algn": "ctr"})
        fill = ln.makeelement(qn("a:solidFill"), {})
        clr = ln.makeelement(qn("a:srgbClr"), {"val": color})
        fill.append(clr)
        ln.append(fill)
        old = tcPr.find(qn(tag[e]))
        if old is not None:
            tcPr.remove(old)
        tcPr.append(ln)


def table(slide, rows, x, y, w, col_w=None, size=12, head_size=12, row_h=0.34,
          align=None, highlight=None):
    """rows[0] = 헤더. PowerPoint 기본 줄무늬 스타일을 끄고 직접 그린다."""
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, In(x), In(y), In(w), In(row_h * nr))
    tbl = gt.table
    tbl.first_row = False
    tbl.horz_banding = False
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(In(w) * cw / total))
    for r, row in enumerate(rows):
        tbl.rows[r].height = In(row_h if r else row_h * 1.05)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = In(0.08)
            cell.margin_right = In(0.06)
            cell.margin_top = cell.margin_bottom = In(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if r == 0 else WHITE
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = (align[c] if align else
                           (PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER))
            run = p.add_run()
            run.text = str(val)
            hot = highlight is not None and r in highlight and r != 0
            set_font(run, head_size if r == 0 else size,
                     bold=(r == 0 or hot),
                     color=MUTED if r == 0 else (ORANGE if hot else INK))
            cell_border(cell, ("B",), "8E8E8B" if r == 0 else "E6E5E2",
                        12700 if r == 0 else 9525)
    return tbl


# ----------------------------------------------------------------- 슬라이드 틀
class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = In(W), In(H)
        self.blank = self.prs.slide_layouts[6]
        self.n = 0

    def _new(self):
        return self.prs.slides.add_slide(self.blank)

    def page(self, slide, foot=None):
        self.n += 1
        tf = textbox(slide, W - MR - 1.4, H - 0.52, 1.4, 0.3)
        para(tf, str(self.n), 10, color=FAINT, align=PP_ALIGN.RIGHT, first=True,
             space_after=0)
        if foot:
            tf2 = textbox(slide, ML, H - 0.52, 8.0, 0.3)
            para(tf2, foot, 9.5, color=FAINT, first=True, space_after=0)

    def head(self, slide, kicker, title, sub=None):
        tf = textbox(slide, ML, TOP_TITLE, W - ML - MR, 0.3)
        para(tf, kicker, 11, bold=True, color=ORANGE, first=True, space_after=3)
        tf2 = textbox(slide, ML, TOP_TITLE + 0.28, W - ML - MR, 0.5)
        para(tf2, title, 25, bold=True, color=INK, first=True, space_after=0)
        y = TOP_TITLE + 0.92
        if sub:
            tf3 = textbox(slide, ML, y, W - ML - MR, 0.34)
            para(tf3, sub, 13, color=MUTED, first=True, space_after=0)
            y += 0.34
        rect(slide, ML, y + 0.08, W - ML - MR, 0.011, fill=RULE)
        return y + 0.34

    def content(self, kicker, title, sub=None, foot=None):
        s = self._new()
        top = self.head(s, kicker, title, sub)
        self.page(s, foot)
        return s, top


# ----------------------------------------------------------------- 슬라이드들
TITLE = "갑상선 결절 초음파 영상의 경량 검출·분할 모델"
SUBTITLE = "추론 비용 없는 경계 정교화 모듈과 INT8 양자화를 중심으로"


def s_title(d):
    s = d._new()
    rect(s, 0, 0, W, H, fill=WHITE)
    rect(s, 0, 0, 0.09, H, fill=ORANGE)
    rect(s, ML, 2.32, 1.5, 0.028, fill=INK)
    tf = textbox(s, ML, 1.72, 10.4, 0.4)
    para(tf, "석사학위논문 발표", 12.5, bold=True, color=MUTED, first=True, space_after=0)
    tf = textbox(s, ML, 2.62, 11.2, 1.4)
    para(tf, TITLE, 36, bold=True, color=INK, first=True, space_after=10, line=1.18)
    para(tf, SUBTITLE, 17, color=MUTED, line=1.3)

    tf = textbox(s, ML, 5.28, 11.2, 1.2)
    para(tf, "YOLOv11-seg · Boundary Refinement · Post-Training Quantization",
         12, color=FAINT, first=True, space_after=14)
    para(tf, "발표자  ○○○          지도교수  ○○○", 13.5, color=INK, space_after=4)
    para(tf, "2026. 09.", 12, color=MUTED)
    d.n += 1
    return s


def s_agenda(d):
    s, top = d.content("Contents", "발표 순서")
    items = [
        ("01", "연구 배경 및 문제 정의",
         "갑상선 초음파 판독에서 경계가 갖는 임상적 의미와 자동화의 한계"),
        ("02", "선행연구 및 연구가설",
         "분할·경계정교화·경량화 세 계열의 정리와 그로부터 도출한 세 가설"),
        ("03", "제안 방법",
         "학습 시에만 동작하는 경계 분기(SegmentBRM)와 경계 가중 손실, 양자화 파이프라인"),
        ("04", "실험 및 결과",
         "5-fold 교차검증 · held-out 평가 · 통계 검정 · 정밀도별 경량화 비교"),
    ]
    y = top + 0.24
    for num, t, sub in items:
        rect(s, ML, y + 0.06, 0.55, 0.55, fill=PANEL)
        tf = textbox(s, ML, y + 0.15, 0.55, 0.4)
        para(tf, num, 15, bold=True, color=ORANGE, align=PP_ALIGN.CENTER, first=True,
             space_after=0)
        tf = textbox(s, ML + 0.82, y + 0.04, 10.4, 0.7)
        para(tf, t, 18, bold=True, first=True, space_after=3)
        para(tf, sub, 12.5, color=MUTED)
        y += 1.18
    return s


def s_section(d, num, title, lead):
    s = d._new()
    rect(s, 0, 0, W, H, fill=WHITE)
    rect(s, 0, 0, 0.09, H, fill=ORANGE)
    tf = textbox(s, ML, 2.75, 10.6, 0.5)
    para(tf, num, 15, bold=True, color=ORANGE, first=True, space_after=8)
    para(tf, title, 34, bold=True, color=INK, space_after=12)
    rect(s, ML, 4.42, 1.2, 0.024, fill=RULE)
    tf = textbox(s, ML, 4.68, 9.6, 0.5)
    para(tf, lead, 14, color=MUTED, first=True, space_after=0, line=1.35)
    d.page(s)
    return s


def s_background(d):
    s, top = d.content(
        "01  연구 배경 및 문제 정의", "초음파 판독에서 '경계'는 부수적 정보가 아니다",
        foot="TI-RADS: Thyroid Imaging Reporting and Data System")
    tf = textbox(s, ML, top + 0.18, 5.55, 4.6)
    para(tf, "갑상선 결절은 성인 초음파 검사에서 흔하게 발견되며, 대부분은 양성이지만 "
             "악성 감별을 위해 매 검사마다 형태 평가가 필요하다.",
         13.5, first=True, space_after=13, line=1.36)
    para(tf, "ACR TI-RADS를 포함한 주요 판독 체계에서 결절의 경계(margin)는 "
             "구성·에코·형태와 나란히 놓이는 독립 채점 항목이다. 매끄러운 경계와 "
             "분엽상·불규칙 경계는 서로 다른 점수를 받고, 이 점수가 세침흡인 여부를 좌우한다.",
         13.5, space_after=13, line=1.36)
    para(tf, "따라서 자동 분할에 요구되는 것은 '영역이 대체로 맞는 것'이 아니라 "
             "'경계가 정확한 것'이다. 그러나 초음파 영상은 스페클 잡음·후방 음영·낮은 대조도로 "
             "경계 자체가 모호하고, 판독자 간 변이도 크다.",
         13.5, space_after=0, line=1.36)

    rect(s, 6.55, top + 0.12, 6.1, 0.028, fill=ORANGE)
    picture(s, f"{FIG}/fig10_dataset.png", 6.55, top + 0.24, 6.1, 3.3)
    tf = textbox(s, 6.55, top + 3.66, 6.1, 0.8)
    para(tf, "본 연구에 사용한 두 공개 데이터셋의 표본과 정답 경계(초록). "
             "장비·화각·주석 방식이 서로 다르다는 점이 이후 교차 데이터셋 실험의 출발점이 된다.",
         11, color=MUTED, first=True, space_after=0, line=1.3)
    return s


def s_problem(d):
    s, top = d.content("01  연구 배경 및 문제 정의", "세 가지 문제 정의",
                       "각 문제는 이후 Phase 2~4의 실험 설계에 그대로 대응된다")
    items = [
        ("P1", "경계 품질이 지표에 가려진다",
         "mAP·IoU는 영역 중심 지표라 몇 픽셀 폭의 경계 오차에 둔감하다. "
         "경계가 실제로 나아졌는지 판정하려면 Boundary IoU·HD95 같은 경계 특화 지표가 "
         "함께 필요하다.", BLUE),
        ("P2", "데이터셋이 서로 이질적이다",
         "공개 데이터셋은 장비·화각·주석 스타일이 달라, 한 데이터셋에서 학습한 모델이 "
         "다른 데이터셋에서도 통할지 알 수 없다. 혼합 학습과 교차 데이터셋 검증이 필요하다.",
         ORANGE),
        ("P3", "배포 환경에 GPU가 없다",
         "검진 현장·휴대형 초음파에는 GPU가 없는 경우가 많다. 정확도를 지키면서 "
         "CPU 실시간 추론이 가능한 수준까지 모델을 줄여야 실제로 쓰인다.", BLUE),
    ]
    x = ML
    cw = (W - ML - MR - 0.5) / 3
    for tag, t, body, col in items:
        rect(s, x, top + 0.2, cw, 2.95, fill=PANEL)
        rect(s, x, top + 0.2, cw, 0.045, fill=col)
        tf = textbox(s, x + 0.3, top + 0.52, cw - 0.6, 3.2)
        para(tf, tag, 13, bold=True, color=col, first=True, space_after=7)
        para(tf, t, 16.5, bold=True, space_after=11, line=1.25)
        para(tf, body, 12.5, color=MUTED, space_after=0, line=1.4)
        x += cw + 0.25
    tf = textbox(s, ML, top + 3.42, W - ML - MR, 0.4)
    para(tf, "→ 본 연구는 세 문제를 순차적으로 다룬다: 데이터 혼합·분할(P2) → "
             "경계 정교화(P1) → 양자화 경량화(P3).", 13, bold=True, first=True, space_after=0)
    return s


def s_related(d):
    s, top = d.content("02  선행연구 및 연구가설", "선행연구 정리와 연구 공백")
    rows = [
        ["계열", "대표 연구", "본 연구와의 관계"],
        ["의료영상 분할", "U-Net (Ronneberger et al., 2015) 및 파생 구조",
         "정확도는 높으나 검출·분할 통합과 추론 비용은 주된 관심이 아님"],
        ["갑상선 결절 특화", "TRFE-Net · TN3K 데이터셋 (Gong et al., 2021·2023)",
         "갑상선 영역 사전정보를 활용. 경량화·배포는 범위 밖"],
        ["실시간 인스턴스 분할", "YOLACT, YOLOv8-seg, YOLOv11-seg (Ultralytics)",
         "프로토타입 마스크로 빠르지만 마스크 해상도가 낮아 경계가 거칠다"],
        ["경계 정교화", "PointRend (Kirillov et al., 2020), Boundary loss "
         "(Kervadec et al., 2019), HD loss (Karimi & Salcudean, 2020)",
         "추론 시 추가 연산이 필요하거나 semantic segmentation을 전제로 함"],
        ["경계 평가 지표", "Boundary IoU (Cheng et al., 2021)",
         "본 연구의 주 평가지표로 채택"],
    ]
    table(s, rows, ML, top + 0.2, W - ML - MR, col_w=[2.0, 4.3, 5.0],
          size=12, head_size=11.5, row_h=0.62,
          align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    y = top + 4.05
    rect(s, ML, y, W - ML - MR, 0.82, fill=PANEL)
    rect(s, ML, y, 0.045, 0.82, fill=ORANGE)
    tf = textbox(s, ML + 0.3, y + 0.15, W - ML - MR - 0.6, 0.6)
    para(tf, "연구 공백", 11.5, bold=True, color=ORANGE, first=True, space_after=5)
    para(tf, "갑상선 초음파 인스턴스 분할에서, 추론 비용을 늘리지 않는 경계 정교화와 "
             "실제 배포를 위한 양자화를 하나의 파이프라인에서 함께 검증한 연구는 드물다.",
         13, space_after=0)
    return s


def s_hypothesis(d):
    s, top = d.content("02  선행연구 및 연구가설", "연구가설",
                       "각 가설은 검증 방법과 기각 조건을 함께 정의했다")
    items = [
        ("H1", "모델 용량 가설",
         "YOLOv11-seg nano는 small 대비 파라미터가 1/3.5 수준임에도 "
         "갑상선 결절 검출·분할에서 동등한 정확도를 낸다.",
         "검증  pooled 5-fold + held-out test에서 nano·small의 Mask mAP50-95 비교"),
        ("H2", "경계 정교화 가설",
         "학습 시에만 동작하는 보조 경계 분기와 경계 가중 손실은 "
         "추론 비용을 전혀 늘리지 않고 경계 품질을 개선한다.",
         "검증  Boundary IoU·HD95·Dice의 대응표본 검정 + GFLOPs 동일성 확인"),
        ("H3", "경량화 가설",
         "INT8 사후 양자화로 CPU 배포 요건을 만족시키면서 "
         "정확도 손실은 무시할 수 있는 수준으로 유지된다.",
         "검증  5개 정밀도(PyTorch·ONNX FP32/FP16/INT8·OpenVINO INT8) 정확도·지연 비교"),
    ]
    y = top + 0.22
    for tag, name, body, how in items:
        rect(s, ML, y, W - ML - MR, 1.44, fill=WHITE, line=RULE)
        rect(s, ML, y, 0.05, 1.44, fill=ORANGE if tag == "H2" else BLUE)
        tf = textbox(s, ML + 0.34, y + 0.17, 1.5, 0.5)
        para(tf, tag, 21, bold=True, color=ORANGE if tag == "H2" else BLUE,
             first=True, space_after=2)
        para(tf, name, 10.5, color=MUTED)
        tf = textbox(s, ML + 2.0, y + 0.2, 9.6, 1.1)
        para(tf, body, 14.5, first=True, space_after=8, line=1.3)
        para(tf, how, 11.5, color=MUTED, space_after=0)
        y += 1.62
    return s


def s_pipeline(d):
    s, top = d.content("03  제안 방법", "전체 연구 파이프라인",
                       "데이터 정비부터 배포 형식 선택까지를 하나의 흐름으로 설계했다")
    picture(s, f"{FIG}/fig8_pipeline.png", ML, top + 0.02, W - ML - MR, 4.72)
    return s


def s_data_intro(d):
    s, top = d.content("03  제안 방법", "사용한 두 공개 데이터셋",
                       "서로 다른 기관·장비에서 만들어진 데이터셋을 의도적으로 함께 사용했다")
    cols = [
        (ML, "TN3K", BLUE, "Gong et al., 2021 · 2023 공개",
         f"{FIG}/fig11_tn3k_samples.png",
         [("규모", "3,493장 · 결절 3,820개 (trainval 2,879 / test 614)"),
          ("형식", "JPG · 184×174 ~ 1463×826, 해상도가 2,443종으로 제각각"),
          ("촬영 특성", "병변 주변을 잘라낸 확대 영상. 측정 커서(캘리퍼)가 "
                     "화면에 남아 있는 경우가 있다"),
          ("결절 수", "이미지당 평균 1.09개 — 7.8%의 영상에 결절이 2개 이상")]),
        (6.85, "DDTI", ORANGE, "Pedraza et al., 2015 · 콜롬비아 IDIME 영상센터",
         f"{FIG}/fig12_ddti_samples.png",
         [("규모", "637장 · 결절 637개 (전량을 학습·검증에 사용)"),
          ("형식", "PNG · 높이는 항상 360 px, 너비는 280·560·315·245의 4종"),
          ("촬영 특성", "장비 화면을 그대로 캡처. 게인 막대·프로브 정보·좌우 표기가 함께 찍힌다"),
          ("결절 수", "모든 영상이 정확히 1개 — 다중 결절 사례가 없다")]),
    ]
    cw = 5.8
    for x, name, color, src, figpath, facts in cols:
        tf = textbox(s, x, top + 0.08, cw, 0.6)
        para(tf, name, 20, bold=True, color=color, first=True, space_after=2)
        para(tf, src, 11, color=MUTED, space_after=0)
        picture(s, figpath, x, top + 0.82, cw, 1.30)
        rect(s, x, top + 2.32, cw, 0.012, fill=RULE)
        tf = textbox(s, x, top + 2.50, cw, 2.3)
        for i, (k, v) in enumerate(facts):
            para(tf, k, 11.5, bold=True, color=color, first=(i == 0),
                 space_before=0 if i == 0 else 9, space_after=2)
            para(tf, v, 12, color=INK, space_after=0, line=1.34)
    return s


def s_data_gap(d):
    s, top = d.content("03  제안 방법", "두 데이터셋의 정량적 차이",
                       "이 격차가 교차 데이터셋(LODO) 실험을 설계한 이유이자, "
                       "그 결과를 해석하는 근거가 된다")
    picture(s, f"{FIG}/fig13_dataset_stats.png", ML, top + 0.08, W - ML - MR, 3.05)
    y = top + 3.35
    items = [
        ("해상도", "TN3K는 2,443종의 서로 다른 크기, DDTI는 4종 고정. 학습 시 640×640 "
                 "레터박스로 통일되지만 원본 스케일 분포는 크게 다르다.", BLUE),
        ("결절 크기", "TN3K는 화면의 0.06%~43%까지 극단이 넓고, DDTI는 1.5%~24%에 "
                   "몰려 있다. DDTI만 본 모델은 아주 작은 결절을 학습한 적이 없다.", ORANGE),
        ("결절 개수", "DDTI는 100%가 단일 결절이다. 다중 결절 장면을 한 번도 학습하지 "
                   "못하므로, TN3K로 옮겨갈 때 놓치는 병변이 생긴다.", ORANGE),
    ]
    for i, (t, b, col) in enumerate(items):
        x = ML + i * ((W - ML - MR) / 3)
        tf = textbox(s, x, y, (W - ML - MR) / 3 - 0.32, 1.2)
        para(tf, t, 12.5, bold=True, color=col, first=True, space_after=4)
        para(tf, b, 11.5, color=MUTED, space_after=0, line=1.34)
    tf = textbox(s, ML, y + 1.32, W - ML - MR, 0.4)
    para(tf, "→ LODO에서 DDTI→TN3K 방향의 mAP가 0.307로 낮은 것은 모델의 결함이라기보다 "
             "학습 데이터가 대상 분포를 포함하지 못한 결과로 읽어야 한다.",
         12.5, bold=True, first=True, space_after=0, line=1.3)
    return s


def s_data_split(d):
    s, top = d.content("03  제안 방법", "전처리와 분할 설계",
                       "테스트 집합은 전 실험에서 단 한 번도 학습에 사용하지 않았다")
    rows = [
        ["구분", "출처", "장수", "결절 수", "용도"],
        ["TN3K (trainval)", "TN3K", "2,879", "3,130", "학습·검증"],
        ["DDTI", "DDTI", "637", "637", "학습·검증"],
        ["TN3K (test)", "TN3K", "614", "690", "최종 평가 전용"],
        ["합계", "", "4,130", "4,457", ""],
    ]
    table(s, rows, ML, top + 0.18, 6.4, col_w=[2.1, 1.0, 1.0, 1.1, 1.7], size=12,
          row_h=0.42, highlight={4})

    tf = textbox(s, ML, top + 2.58, 6.4, 2.3)
    para(tf, "전처리", 13, bold=True, color=ORANGE, first=True, space_after=8)
    para(tf, "두 데이터셋의 정답이 비트맵 마스크로 제공되어 YOLO-seg 폴리곤으로 "
             "변환했다. 임계값 127 이진화 → 모폴로지 open으로 잡음 제거 → 외곽 윤곽만 "
             "추출하는 절차를 두 데이터셋에 동일하게 적용했다.",
         12, space_after=8, line=1.35)
    para(tf, "변환 후 4,457개 폴리곤 전수를 원본 마스크와 겹쳐 그려 대조했고, 빈 라벨이나 "
             "마스크 누락이 한 건도 없음을 확인했다.",
         12, color=MUTED, space_after=0, line=1.35)

    tf = textbox(s, 7.35, top + 0.14, W - MR - 7.35, 4.5)
    para(tf, "분할 설계", 13, bold=True, color=ORANGE, first=True, space_after=9)
    for t, b in [
        ("Pooled 5-fold", "TN3K trainval + DDTI 3,516장을 출처 계층화(source-stratified)로 "
                          "5분할. 두 데이터셋의 비율이 각 fold에서 동일하게 유지된다. "
                          "fold당 학습 2,812장 / 검증 704장."),
        ("Held-out test", "TN3K test 614장. Phase 2~5의 공통 최종 평가 집합이며, "
                          "모델 선택에도 사용하지 않았다."),
        ("LODO (Leave-One-Dataset-Out)",
         "TN3K→DDTI, DDTI→TN3K 양방향. 학습에 쓰지 않은 데이터셋에서 평가해 "
         "도메인 일반화를 직접 측정한다. 시드 42·43·44로 각 3회 반복했다."),
    ]:
        para(tf, t, 13, bold=True, space_before=4, space_after=3)
        para(tf, b, 12, color=MUTED, space_after=9, line=1.35)
    return s


def s_baseline(d):
    s, top = d.content("03  제안 방법", "베이스라인 선정: nano인가 small인가",
                       "H1 검증 — 두 모델을 동일 조건으로 5-fold 학습 후 held-out에서 비교")
    rows = [
        ["", "YOLOv11n-seg (nano)", "YOLOv11s-seg (small)"],
        ["Mask mAP50-95 (test)", "0.6272", "0.6239"],
        ["Box mAP50-95 (test)", "0.6216", "0.6181"],
        ["파라미터", "2.84 M", "10.08 M"],
        ["GFLOPs", "9.73", "33.08"],
        ["가중치 크기", "6.00 MB", "20.51 MB"],
    ]
    table(s, rows, ML, top + 0.2, 7.4, col_w=[2.9, 2.3, 2.3], size=13, row_h=0.45,
          highlight={1})
    tf = textbox(s, 8.5, top + 0.24, W - MR - 8.5, 3.6)
    para(tf, "결론", 13, bold=True, color=ORANGE, first=True, space_after=8)
    para(tf, "정확도는 nano가 오히려 근소하게 앞선다. 파라미터 3.5배, 연산량 3.4배를 "
             "더 쓰고도 small이 얻는 것이 없다.", 13.5, space_after=10, line=1.35)
    para(tf, "결절이 단일 클래스이고 형태 변이가 제한적이어서, small의 추가 용량이 "
             "이 과제에서는 활용되지 않은 것으로 해석된다.", 12.5, color=MUTED,
         space_after=10, line=1.35)
    para(tf, "→ 이후 모든 실험은 nano를 기준 모델로 진행한다. H1 지지.",
         13, bold=True, space_after=0, line=1.3)
    return s


def s_arch(d):
    s, top = d.content("03  제안 방법", "제안 구조: SegmentBRM 헤드",
                       "Proto 특징에서 갈라지는 보조 경계 분기를 두되, 추론 그래프에는 남기지 않는다")
    picture(s, f"{FIG}/fig5_architecture.png", ML, top + 0.02, 7.5, 4.7)
    tf = textbox(s, 8.35, top + 0.3, W - MR - 8.35, 4.2)
    para(tf, "설계 원칙", 13, bold=True, color=ORANGE, first=True, space_after=9)
    for t, b in [
        ("기존 헤드를 대체하지 않고 확장",
         "Detect·마스크 계수(cv4)·Proto 세 경로는 그대로 두고, Proto 출력에서 "
         "경계 분기만 추가로 갈라낸다."),
        ("분기는 최소 크기",
         "Conv(32→16, 3×3) + Conv(16→1, 1×1), 총 4,657개 파라미터. "
         "전체의 0.16%에 불과하다."),
        ("추론에서는 실행하지 않음",
         "학습 시에만 4-튜플을 반환하고, 추론 모드에서는 분기 자체를 호출하지 않는다. "
         "따라서 GFLOPs는 베이스라인과 정확히 같은 9.73이다."),
    ]:
        para(tf, t, 12.5, bold=True, space_before=3, space_after=3)
        para(tf, b, 11.5, color=MUTED, space_after=9, line=1.34)
    return s


def s_dual(d):
    s, top = d.content("03  제안 방법", "학습 경로와 추론 경로의 분리",
                       "경계 감독은 학습에만 개입하고, 배포되는 그래프는 베이스라인과 동일하다")
    picture(s, f"{FIG}/fig6_dual_path.png", ML, top + 0.15, W - ML - MR, 4.0)
    tf = textbox(s, ML, top + 4.32, W - ML - MR, 0.5)
    para(tf, "이 분리 덕분에 '정확도 향상에는 대가가 따른다'는 통상적 절충을 피할 수 있다. "
             "경계 분기가 학습 중 Proto 특징을 경계에 민감하도록 밀어주고, 그 효과는 "
             "가중치에 남은 채 분기만 제거된다.", 12.5, color=MUTED, first=True,
         space_after=0, line=1.35)
    return s


def s_loss(d):
    s, top = d.content("03  제안 방법", "경계 가중 마스크 손실",
                       "정답 마스크의 형태학적 경계 밴드에 손실 가중치를 집중시킨다")
    picture(s, f"{FIG}/fig7_boundary_loss.png", ML, top + 0.1, W - ML - MR, 3.35)
    y = top + 3.62
    for i, (t, b) in enumerate([
        ("경계 밴드 추출", "B = dilate(M) − erode(M). 3×3 max-pool 두 번으로 구현되어 "
                       "GPU에서 추가 비용이 거의 없다."),
        ("가중 BCE", "w = 1 + λ·B (λ=2.0). 경계 부근 픽셀의 오차가 내부 픽셀보다 "
                   "3배 크게 반영된다."),
        ("보조 경계 손실", "경계 분기 출력과 정답 경계 사이의 BCE를 head_gain=1.0으로 "
                      "더한다. 두 계수는 스윕으로 결정했다."),
    ]):
        x = ML + i * ((W - ML - MR) / 3)
        tf = textbox(s, x, y, (W - ML - MR) / 3 - 0.3, 1.0)
        para(tf, t, 12.5, bold=True, color=ORANGE, first=True, space_after=4)
        para(tf, b, 11.5, color=MUTED, space_after=0, line=1.34)
    return s


def s_quant(d):
    s, top = d.content("03  제안 방법", "경량화: 사후 양자화 파이프라인",
                       "재학습 없이 배포 형식만 바꿔 크기와 CPU 지연을 줄인다")
    tf = textbox(s, ML, top + 0.18, 5.9, 4.2)
    para(tf, "구성", 13, bold=True, color=ORANGE, first=True, space_after=9)
    for t, b in [
        ("ONNX FP32 / FP16", "ultralytics export, opset 13, simplify 적용."),
        ("onnxruntime 정적 PTQ",
         "QDQ 형식, per-channel, 가중치 INT8·활성값 UINT8. "
         "TN3K trainval에서 150장을 무작위 추출해 캘리브레이션."),
        ("OpenVINO INT8 (NNCF)",
         "YOLO 구조를 인식해 민감 노드를 자동 제외하는 PTQ."),
    ]:
        para(tf, t, 13, bold=True, space_before=3, space_after=3)
        para(tf, b, 12, color=MUTED, space_after=9, line=1.35)

    tf = textbox(s, 7.1, top + 0.14, W - MR - 7.1, 4.4)
    para(tf, "구현 과정에서 확인한 함정", 13, bold=True, color=ORANGE, first=True,
         space_after=9)
    for t, b in [
        ("검출 헤드를 양자화하면 모델이 붕괴한다",
         "INT8화 시 분류 점수가 전부 0이 되었다. /model.23/ 이하 119개 노드를 "
         "양자화에서 제외해 해결했다."),
        ("per-channel QDQ는 opset 13 이상을 요구한다",
         "opset 12에서는 InvalidGraph 오류로 로드되지 않는다."),
        ("OpenVINO 백엔드는 프로세스의 CUDA를 비활성화한다",
         "평가 순서를 정밀도 단위로 재구성하고 OpenVINO를 마지막에 두어 회피했다."),
    ]:
        para(tf, t, 12.5, bold=True, space_before=3, space_after=3)
        para(tf, b, 11.5, color=MUTED, space_after=9, line=1.34)

    y = top + 3.62
    rect(s, ML, y, W - ML - MR, 1.08, fill=PANEL)
    rect(s, ML, y, 0.045, 1.08, fill=ORANGE)
    tf = textbox(s, ML + 0.3, y + 0.2, W - ML - MR - 0.6, 0.8)
    para(tf, "세 함정 모두 조용히 실패한다", 11.5, bold=True, color=ORANGE, first=True,
         space_after=5)
    para(tf, "오류 메시지 없이 정확도만 0에 가깝게 떨어지거나 프로세스가 종료된다. "
             "양자화 결과는 반드시 원본과 같은 평가 집합에서 정확도를 재측정해 확인해야 한다.",
         13, space_after=0, line=1.32)
    return s


def s_setup(d):
    s, top = d.content("04  실험 및 결과", "실험 설정",
                       "베이스라인과 제안 모델은 학습 하이퍼파라미터를 완전히 동일하게 두었다")
    rows = [
        ["항목", "설정"],
        ["학습", "100 epoch (patience 20), imgsz 640, batch 32, seed 42, deterministic"],
        ["마스크", "overlap_mask, mask_ratio 4"],
        ["평가 집합", "held-out TN3K test 614장 (전 Phase 공통)"],
        ["추론 설정", "conf 0.25, IoU 0.7, retina_masks"],
        ["정확도 지표", "Mask mAP50-95, Dice, Boundary IoU (d = 0.02 × 대각선), "
                   "HD95 (대각선 대비 %)"],
        ["효율 지표", "파라미터, GFLOPs, 파일 크기, CPU·GPU 지연 (batch 1, "
                  "워밍업 15회 후 100회 평균)"],
        ["하드웨어", "NVIDIA A100 80GB × 3 (학습·GPU 지연), 동일 서버 CPU (CPU 지연)"],
    ]
    table(s, rows, ML, top + 0.2, W - ML - MR, col_w=[2.4, 9.6], size=12.5,
          row_h=0.55, align=[PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    return s


def s_r_boundary(d):
    s, top = d.content("04  실험 및 결과", "결과 ① 경계 정교화 효과",
                       "held-out test, 5-fold 각각을 짝지어 비교")
    picture(s, f"{FIG}/fig1_boundary_improvement.png", ML, top + 0.1, 8.0, 4.3)
    tf = textbox(s, 8.85, top + 0.35, W - MR - 8.85, 4.0)
    for t, b, col in [
        ("Dice  +0.0104", "0.8195 → 0.8300. 5개 fold 전부에서 개선.", ORANGE),
        ("HD95  −13.0%", "9.628 → 8.378 (대각선 대비 %). 경계 최악 오차가 줄었다.", ORANGE),
        ("Boundary IoU  +0.0055", "0.4358 → 0.4414.", ORANGE),
        ("mAP50-95  −0.0005", "0.6272 → 0.6267. 검출 성능은 그대로 유지된다.", MUTED),
        ("GFLOPs  ±0", "9.73으로 완전히 동일. 추론 비용 증가가 없다.", BLUE),
    ]:
        para(tf, t, 15, bold=True, color=col, first=(col is ORANGE and "Dice" in t),
             space_before=6, space_after=3)
        para(tf, b, 11.5, color=MUTED, space_after=8, line=1.34)
    return s


def s_r_qual(d):
    s, top = d.content("04  실험 및 결과", "결과 ② 정성 비교",
                       "개선 폭 분포의 분위수에서 뽑아 실패 사례까지 함께 제시한다")
    picture(s, f"{FIG}/fig4_qualitative.png", ML, top + 0.02, 4.6, 4.7)
    tf = textbox(s, 5.75, top + 0.35, W - MR - 5.75, 4.2)
    para(tf, "표본 선택 방식", 13, bold=True, color=ORANGE, first=True, space_after=8)
    para(tf, "개선 폭이 큰 사례만 고르면 그림이 결론을 앞서간다. 그래서 Dice 개선 폭으로 "
             "정렬한 뒤 최상위·상위 4분위·중앙값·최하위에서 한 장씩 뽑았다.",
         13, space_after=12, line=1.35)
    para(tf, "읽는 법", 13, bold=True, color=ORANGE, space_after=8)
    para(tf, "왼쪽부터 정답 경계, 베이스라인 예측, 제안 모델 예측이다. 상단 세 행에서는 "
             "제안 모델의 윤곽이 정답에 더 밀착하고, 특히 결절 하단의 음영 구간에서 "
             "베이스라인이 새어 나가는 부분이 억제된다.", 13, space_after=12, line=1.35)
    para(tf, "마지막 행은 제안 모델이 더 나빠진 사례다. 배경 조직을 결절로 잘못 포함해 "
             "Dice가 크게 떨어졌다. 경계 가중이 오히려 잘못된 경계를 강화할 수 있음을 "
             "보여주는 실패 유형으로, 후속 과제로 남긴다.", 12.5, color=MUTED,
         space_after=0, line=1.35)
    return s


def s_r_stats(d):
    s, top = d.content("04  실험 및 결과", "결과 ③ 통계적 유의성과 교차 데이터셋 검증",
                       "평균 비교를 넘어, 개선이 우연의 범위인지 검정했다")
    picture(s, f"{FIG}/fig9_significance.png", ML, top + 0.2, W - ML - MR, 3.15)
    y = top + 3.62
    for i, (t, b, col) in enumerate([
        ("in-domain: 세 지표 모두 유의",
         "이미지 단위 대응표본(n=614)에서 Dice p=0.014, Boundary IoU p=0.011, "
         "HD95 p=0.0045. 세 신뢰구간 모두 0을 포함하지 않는다.", ORANGE),
        ("cross-domain: 판정 불가",
         "시드를 3개(42·43·44)로 늘리자 차이가 크게 줄었고, 시드 단위 검정에서는 "
         "전 지표가 유의하지 않다(p = 0.51 ~ 1.00).", BLUE),
        ("단일 시드의 위험",
         "시드 42만 보면 Dice −0.118로 '교차 도메인 파탄'처럼 보이나, 이는 해당 "
         "실행 하나의 이상치였다. 실행 간 변동이 효과보다 컸다.", MUTED),
    ]):
        x = ML + i * ((W - ML - MR) / 3)
        tf = textbox(s, x, y, (W - ML - MR) / 3 - 0.3, 1.1)
        para(tf, t, 12.5, bold=True, color=col, first=True, space_after=4)
        para(tf, b, 11.5, color=MUTED, space_after=0, line=1.34)
    return s


def s_r_quant(d):
    s, top = d.content("04  실험 및 결과", "결과 ④ 정밀도별 경량화 비교",
                       "5-fold 평균, held-out test 기준")
    picture(s, f"{FIG}/fig3_quantization.png", ML, top + 0.04, W - ML - MR, 2.42)
    rows = [
        ["정밀도", "크기(MB)", "CPU(ms)", "GPU(ms)", "mAP50-95", "Dice", "HD95%"],
        ["PyTorch FP32", "6.01", "–", "–", "0.6266", "0.8300", "8.378"],
        ["ONNX FP32", "11.63", "62.82", "3.53", "0.6188", "0.8300", "8.529"],
        ["ONNX FP16", "5.87", "–", "3.13", "0.6194", "0.8301", "8.516"],
        ["ONNX INT8", "5.45", "44.00", "–", "0.6032", "0.8184", "9.395"],
        ["OpenVINO INT8", "3.71", "4.10", "–", "0.6161", "0.8290", "8.497"],
    ]
    table(s, rows, ML, top + 2.66, W - ML - MR, col_w=[2.4, 1.5, 1.5, 1.5, 1.7, 1.5, 1.5],
          size=11.5, head_size=11, row_h=0.32, highlight={5})
    return s


def s_r_pareto(d):
    s, top = d.content("04  실험 및 결과", "결과 ⑤ 정확도와 효율의 절충",
                       "CPU 지연을 기준으로 본 배포 후보 비교")
    picture(s, f"{FIG}/fig2_accuracy_vs_latency.png", ML, top + 0.05, 6.4, 4.6)
    tf = textbox(s, 7.35, top + 0.4, W - MR - 7.35, 4.2)
    para(tf, "OpenVINO INT8이 지배적 선택", 15, bold=True, color=ORANGE, first=True,
         space_after=9)
    para(tf, "CPU 지연 62.82 ms → 4.10 ms로 15.3배 빨라지고, 파일 크기는 3.71 MB로 "
             "가장 작다. 그러면서 Dice는 0.8290으로 원본(0.8300)과 사실상 같다.",
         13, space_after=12, line=1.35)
    para(tf, "같은 INT8이라도 onnxruntime 정적 PTQ는 44.00 ms에 그쳤다. 검출 헤드를 "
             "float으로 남겨야 했기 때문이며, NNCF가 민감 노드를 더 정교하게 "
             "선별한다는 점이 차이를 만들었다.", 12.5, color=MUTED, space_after=12,
         line=1.35)
    para(tf, "권장 배포 구성", 13, bold=True, space_after=5)
    para(tf, "CPU 환경 — OpenVINO INT8 (3.71 MB, 4.10 ms)", 12.5, space_after=3)
    para(tf, "GPU 환경 — ONNX FP16 (5.87 MB, 3.13 ms, 무손실)", 12.5, space_after=0)
    return s


def s_summary(d):
    s, top = d.content("04  실험 및 결과", "Phase 2~4 통합 결과",
                       "베이스라인에서 최종 배포 모델까지의 누적 변화")
    rows = [
        ["모델", "mAP50-95", "Dice", "HD95%", "BIoU", "파라미터", "GFLOPs",
         "크기(MB)", "CPU(ms)"],
        ["베이스라인 nano (Phase 2)", "0.6272", "0.8195", "9.628", "0.4358",
         "2.84 M", "9.73", "6.00", "–"],
        ["+ 경계 정교화 (Phase 3)", "0.6267", "0.8300", "8.378", "0.4414",
         "2.85 M", "9.73", "6.01", "–"],
        ["+ ONNX FP16", "0.6194", "0.8301", "8.516", "0.4403", "–", "–", "5.87", "–"],
        ["+ ONNX INT8", "0.6032", "0.8184", "9.395", "0.4299", "–", "–", "5.45", "44.00"],
        ["+ OpenVINO INT8 (최종)", "0.6161", "0.8290", "8.497", "0.4379",
         "–", "–", "3.71", "4.10"],
    ]
    table(s, rows, ML, top + 0.25, W - ML - MR,
          col_w=[3.1, 1.35, 1.15, 1.15, 1.15, 1.25, 1.1, 1.2, 1.2],
          size=11.5, head_size=11, row_h=0.5, highlight={5})
    tf = textbox(s, ML, top + 3.5, W - ML - MR, 0.9)
    para(tf, "경계 지표는 Phase 3에서 개선된 뒤 양자화를 거쳐도 유지된다. "
             "최종 모델은 베이스라인 대비 Dice +0.0095, HD95 −11.7%를 유지하면서 "
             "파일 크기는 1/1.6, CPU 지연은 FP32 대비 1/15.3이다.",
         13, first=True, space_after=0, line=1.4)
    return s


def s_limit(d):
    s, top = d.content("04  실험 및 결과", "한계와 향후 과제",
                       "확인된 것과 확인되지 않은 것을 구분해 기술한다")
    left = [
        ("효과 크기는 작다",
         "경계 개선의 대응표본 효과 크기는 dz ≈ 0.14로 작다. 통계적으로는 유의하지만 "
         "임상적으로 의미 있는 차이인지는 별도 검증이 필요하다."),
        ("교차 데이터셋 일반화는 미확인",
         "3시드 반복 결과 개선도 악화도 유의하지 않았다. '일반화된다'가 아니라 "
         "'현재 표본으로는 판정할 수 없다'가 정확한 서술이다."),
        ("LODO 설정의 낙관적 편향",
         "LODO에서 검증 집합과 평가 집합이 같은 타깃 도메인이므로, 두 모델 모두 "
         "타깃에서 체크포인트를 선택했다. 비교는 공정하나 절대값은 낙관적이다."),
    ]
    right = [
        ("실패 유형 분석",
         "정성 비교의 최악 사례처럼 배경 조직을 결절로 포함하는 오검출을 억제하는 "
         "후처리 또는 손실 설계."),
        ("도메인 적응",
         "장비·주석 스타일 차이를 흡수하는 스타일 정규화나 도메인 적대 학습을 결합해 "
         "교차 데이터셋 성능을 다시 검증."),
        ("엣지 디바이스 실측",
         "Jetson·라즈베리파이 등 실제 배포 후보 장비에서 지연·전력을 측정. "
         "현재 CPU 지연은 서버 CPU 기준값이다."),
    ]
    for col, (items, head, color) in enumerate(
            [(left, "확인된 한계", MUTED), (right, "향후 과제", ORANGE)]):
        x = ML + col * 6.15
        tf = textbox(s, x, top + 0.16, 5.7, 0.3)
        para(tf, head, 13, bold=True, color=color, first=True, space_after=0)
        rect(s, x, top + 0.52, 5.7, 0.014, fill=RULE)
        tf = textbox(s, x, top + 0.7, 5.7, 3.6)
        for i, (t, b) in enumerate(items):
            para(tf, t, 13.5, bold=True, first=(i == 0), space_before=0 if i == 0 else 8,
                 space_after=4)
            para(tf, b, 12, color=MUTED, space_after=10, line=1.36)
    return s


def s_conclusion(d):
    s, top = d.content("Conclusion", "결론")
    items = [
        ("01", "경량 백본에서 출발해도 손해가 없다",
         "nano는 small 대비 파라미터 1/3.5, 연산량 1/3.4임에도 정확도가 근소하게 앞섰다. "
         "갑상선 결절이라는 단일 클래스 과제에서 큰 모델의 추가 용량은 활용되지 않았다."),
        ("02", "경계 정교화를 추론 비용 없이 얻을 수 있다",
         "학습 전용 경계 분기와 경계 가중 손실로 Dice +0.0104, HD95 −13.0%를 얻으면서 "
         "GFLOPs는 9.73으로 동일하다. 세 지표 모두 통계적으로 유의하다."),
        ("03", "CPU 배포는 실용 범위에 들어왔다",
         "OpenVINO INT8로 3.71 MB · CPU 4.10 ms를 달성했고 정확도 손실은 "
         "Dice 0.001 수준이다. GPU 없는 환경에서도 실시간 추론이 가능하다."),
        ("04", "일반화 주장은 유보한다",
         "교차 데이터셋에서는 개선이 유지되지 않았고, 시드 반복 결과 악화 또한 "
         "유의하지 않았다. 단일 실행 결과로 결론을 내리지 않은 점을 성과로 본다."),
    ]
    y = top + 0.2
    for num, t, b in items:
        tf = textbox(s, ML, y, 0.6, 0.4)
        para(tf, num, 13, bold=True, color=ORANGE, first=True, space_after=0)
        tf = textbox(s, ML + 0.72, y - 0.04, 11.2, 0.9)
        para(tf, t, 15.5, bold=True, first=True, space_after=4)
        para(tf, b, 12.5, color=MUTED, space_after=0, line=1.35)
        y += 1.16
    return s


def s_end(d):
    s = d._new()
    rect(s, 0, 0, W, H, fill=WHITE)
    rect(s, 0, 0, 0.09, H, fill=ORANGE)
    tf = textbox(s, ML, 3.1, 10.5, 0.8)
    para(tf, "감사합니다", 34, bold=True, color=INK, first=True, space_after=12)
    para(tf, "질의응답", 15, color=MUTED, space_after=0)
    rect(s, ML, 4.42, 1.2, 0.024, fill=RULE)
    d.n += 1
    return s


def main():
    missing = [f for f in ["fig1_boundary_improvement", "fig2_accuracy_vs_latency",
                           "fig3_quantization", "fig4_qualitative", "fig5_architecture",
                           "fig6_dual_path", "fig7_boundary_loss", "fig8_pipeline",
                           "fig9_significance", "fig10_dataset", "fig11_tn3k_samples",
                           "fig12_ddti_samples", "fig13_dataset_stats"]
               if not os.path.exists(f"{FIG}/{f}.png")]
    assert not missing, f"그림 없음: {missing} — make_figures.py / make_arch_figures.py 먼저 실행"

    d = Deck()
    s_title(d)
    s_agenda(d)

    s_section(d, "01", "연구 배경 및 문제 정의",
              "초음파 판독에서 경계가 갖는 의미와, 자동 분할이 아직 놓치고 있는 지점")
    s_background(d)
    s_problem(d)

    s_section(d, "02", "선행연구 및 연구가설",
              "분할·경계정교화·경량화 세 계열을 정리하고 검증 가능한 형태의 가설로 좁힌다")
    s_related(d)
    s_hypothesis(d)

    s_section(d, "03", "제안 방법",
              "데이터 설계, 학습 전용 경계 분기(SegmentBRM), 경계 가중 손실, 그리고 양자화")
    s_pipeline(d)
    s_data_intro(d)
    s_data_gap(d)
    s_data_split(d)
    s_baseline(d)
    s_arch(d)
    s_dual(d)
    s_loss(d)
    s_quant(d)

    s_section(d, "04", "실험 및 결과",
              "경계 개선 · 정성 비교 · 통계 검정 · 정밀도별 경량화 · 통합 결과")
    s_setup(d)
    s_r_boundary(d)
    s_r_qual(d)
    s_r_stats(d)
    s_r_quant(d)
    s_r_pareto(d)
    s_summary(d)
    s_limit(d)
    s_conclusion(d)
    s_end(d)

    cp = d.prs.core_properties          # 라이브러리 기본 문서 속성 대신 발표 제목만 남긴다
    cp.title, cp.subject = TITLE, SUBTITLE
    cp.author = cp.last_modified_by = cp.comments = cp.keywords = cp.category = ""
    d.prs.save(OUT)
    print(f"저장: {OUT}  ({len(d.prs.slides.__iter__.__self__._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
